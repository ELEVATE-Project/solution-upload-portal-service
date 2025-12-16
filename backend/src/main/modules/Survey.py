import os, uuid, requests, sys, json, time, csv, openpyxl
from dotenv import load_dotenv
from pathlib import Path
from backend.src.main.modules.common_config import *
from openpyxl.styles import Color, PatternFill


env_path = Path(__file__).resolve().parents[1] / "apiServices" / "src" / "main" / ".env"

# Load the .env file
load_dotenv(dotenv_path=env_path)
internal_access_token = os.getenv("internal_access_token")
# adminTokenHeaderName = os.getenv("adminTokenHeaderName")
jwtTokenSecret = os.getenv("jwtTokenSecret")
# projAdminAccessToken = os.getenv("projAdminAccessToken")
# adminAccessToken = os.getenv("adminAccessToken")
authorization = os.getenv("authorization")
authorizationforhost = os.getenv("authorizationforhost")
appname = os.getenv("appname")
x_channel_id = os.getenv("x_channel_id")
host = os.getenv("host")
userLoginHost = os.getenv("userLoginHost")
internal_kong_ip = os.getenv("internal_kong_ip")
elevateprojecthost = os.getenv("elevateprojecthost")
elevateentityhost = os.getenv("elevateentityhost")
identifier = os.getenv("identifier")
password = os.getenv("password")
origin = os.getenv("origin")

class CreateSurvey:
    def __init__(self):
        self.errorVar = []

    def createAPILog(self, solutionName_for_folder_path, messageArr):
        print(solutionName_for_folder_path, "solutionName_for_folder_path")
        file_exists = os.path.join(solutionName_for_folder_path, 'apiHitLogs', 'apiLogs.txt')
        os.makedirs(os.path.dirname(file_exists), exist_ok=True)
        # Create file with header if it doesn't exist
        if not os.path.exists(file_exists):
            with open(file_exists, "w", encoding='utf-8') as API_log:
                API_log.write("===============================================================================\n")
                API_log.write("ENVIRONMENT LOGS\n")
                API_log.write("===============================================================================\n")
        # Append logs
        with open(file_exists, "a", encoding='utf-8') as API_log:
            API_log.write("\n")
            for msg in messageArr:
                API_log.write(str(msg))
                API_log.write("\n")

    def apicheckslog(self, solutionName_for_folder_path, messageArr):
        file_exists = os.path.join(solutionName_for_folder_path, 'apiHitLogs', 'apiLogs.csv')
        fileheader = ["Resource", "Process", "Status", "Remark"]
        os.makedirs(os.path.dirname(file_exists), exist_ok=True)

        # ✅ Create file with header if it doesn't exist
        if not os.path.exists(file_exists):
            with open(file_exists, 'w', newline='', encoding='utf-8') as file:
                writer = csv.writer(file, quoting=csv.QUOTE_NONNUMERIC, delimiter=',', lineterminator='\n')
                writer.writerow(fileheader)  
                
        # ✅ Append a new log entry
        with open(file_exists, 'a', newline='', encoding='utf-8') as file:
            writer = csv.writer(file, quoting=csv.QUOTE_NONNUMERIC, delimiter=',', lineterminator='\n')
            writer.writerow(messageArr)

    def solutionUpdate(self, solutionName_for_folder_path, accessToken, solutionId, bodySolutionUpdate,programdetails):
        try:
            solutionUpdateApi = internal_kong_ip + solutionupdateapi + str(solutionId)
            headerUpdateSolutionApi = {
                'Content-Type': 'application/json',
                'Authorization': authorization,
                'X-auth-token': accessToken,
                'X-Channel-id': x_channel_id,
                "internal-access-token": internal_access_token,
                'tenantId': programdetails.get('TenantID'),
                'orgid': programdetails.get('OrgForAPIs')
                # adminTokenHeaderName: adminAccessToken
                }
            responseUpdateSolutionApi = requests.post(url=solutionUpdateApi, headers=headerUpdateSolutionApi,data=json.dumps(bodySolutionUpdate))
            messageArr = []
            messageArr = ["Solution Update API called.", "URL : " + str(solutionUpdateApi), "Body : " + str(bodySolutionUpdate),"Response : " + str(responseUpdateSolutionApi.text),"Status Code : " + str(responseUpdateSolutionApi.status_code)]
            self.createAPILog(solutionName_for_folder_path, messageArr)
            if responseUpdateSolutionApi.status_code == 200:
                print("Solution Update Success.")
                return True
            else:
                if responseUpdateSolutionApi.status_code in [400, 401, 403, 404, 422]:
                    self.errorVar.append(f"UpdateSolutionApi-Client Error {responseUpdateSolutionApi.status_code}: {responseUpdateSolutionApi.text}")
                elif responseUpdateSolutionApi.status_code in [500, 502, 503, 504]:
                    self.errorVar.append(f"UpdateSolutionApi-Server Error {responseUpdateSolutionApi.status_code}: {responseUpdateSolutionApi.text}")
                else:
                    self.errorVar.append(f"UpdateSolutionApi-Unexpected Error {responseUpdateSolutionApi.status_code}: {responseUpdateSolutionApi.text}")
                # ElevateObservation.createAPILog(solutionName_for_folder_path, errorVar)
                return False
            
        except Exception as e:
            messageArr = []
            messageArr.append("Exception caught : " + str(e))
            self.createAPILog(solutionName_for_folder_path, messageArr)
            self.errorVar.append(f"Error occurred: {str(e)}")
            return False
        
    def createSurveySolution(self, parentFolder, wbObservation, accessToken, programdetails):
        print("Create Survey Solution Func Called....")
        try:
            details = wbObservation.get("details", {})
            if not details:
                raise Exception("❌ 'details' key missing in wbObservation")

            surveySolutionCreationReqBody = {
                "name": details["survey_solution_name"],
                "description": details["survey_solution_description"],
                "externalId": str(uuid.uuid1()),
                "isExternalProgram": True,
                "creator": details["Name_of_the_creator"]
            }

            # userDetails = ElevateObservation.fetchUserDetails(
            #     environment, accessToken,
            #     details["Username/user id/email id/phone no. of the Content creator"]
            # )
            surveySolutionCreationReqBody["author"] = details["Name_of_the_creator"]

            print("🔧 Creating Survey Solution with body:", surveySolutionCreationReqBody)

            urlCreateSolutionApi = internal_kong_ip + surveysolutioncreationapiurl
            headerCreateSolutionApi = {
                'Content-Type': content_type,
                "internal-access-token": internal_access_token,
                'X-auth-token': accessToken,
                'tenantId': programdetails.get('TenantID'),
                'orgid': programdetails.get('OrgForAPIs')
                # adminTokenHeaderName: adminAccessToken
            }
            response = requests.post(
                url=urlCreateSolutionApi,
                headers=headerCreateSolutionApi,
                data=json.dumps(surveySolutionCreationReqBody)
            )
            messageArr = []
            messageArr = ["Survey Solution Creation.",
                        "URL : " + str(urlCreateSolutionApi),
                        "Headers: "+ str(headerCreateSolutionApi),
                        "surveySolutionCreationReqBody: "+str(surveySolutionCreationReqBody),
                        "Status Code : " + str(response.status_code),
                        "Response : " + str(response.text)]
            self.createAPILog(parentFolder, messageArr)
            if response.status_code == 200:
                responseData = response.json()
                solutionId = responseData["result"]["solutionId"]
                print("✅ Solution Created:", solutionId)

                externalIdSearch = surveySolutionCreationReqBody["externalId"]
                urlSearchSolution = f"{internal_kong_ip}{fetchsolutiondetails}survey&page=1&limit=10&search={externalIdSearch}"

                responseSearch = requests.post(urlSearchSolution, headers=headerCreateSolutionApi)
                messageArr = []
                messageArr = ["Survey Solution search Creation.",
                            "URL : " + str(urlSearchSolution),
                            "Headers: "+ str(headerCreateSolutionApi),
                            "Status Code : " + str(responseSearch.status_code),
                            "Response : " + str(responseSearch.text)]
                self.createAPILog(parentFolder, messageArr)
                if responseSearch.status_code == 200:
                    surveySolutionExternalId = responseSearch.json()['result']['data'][0]['externalId']
                    bodySolutionUpdate = {"creator": details["Name_of_the_creator"]}
                    if not self.solutionUpdate(parentFolder, accessToken, solutionId, bodySolutionUpdate,programdetails):
                        self.errorVar.append("Solution Update Failed.")
                        return False
                    return [solutionId, surveySolutionExternalId]

                else:
                    self.errorVar.append(f"❌ Search failed: {responseSearch.text}")
                    return False
            else:
                self.errorVar.append(f"❌ Create solution error: {response.text}")
                return False
        except Exception as e:
            messageArr = []
            messageArr.append("Exception caught : " + str(e))
            self.createAPILog(parentFolder, messageArr)
            self.errorVar.append(f"⚠️ Exception: {str(e)}")
            return False

    def fetchSolutionDetailsFromProgramSheet(self, solutionName_for_folder_path, programdetails, solutionId, accessToken, ProgramGlobalDict):
        try:
            urlFetchSolutionApi = internal_kong_ip + fetchsolutiondoc + solutionId
            headerFetchSolutionApi = {
                'Content-Type': 'application/json',
                'Authorization': authorization,
                'X-auth-token': accessToken,
                'X-Channel-id': x_channel_id,
                'internal-access-token': internal_access_token,
                'tenantId': programdetails.get('TenantID'),
                'orgid': programdetails.get('OrgForAPIs')
                # adminTokenHeaderName: adminAccessToken
            }
            payloadFetchSolutionApi = {}
            responseFetchSolutionApiUrl = requests.post(url=urlFetchSolutionApi, headers=headerFetchSolutionApi,
                                                    data=payloadFetchSolutionApi)
            messageArr = []
            messageArr = ["Solution Fetch Link.",
                        "URL : " + str(urlFetchSolutionApi),
                        "Headers: "+ str(headerFetchSolutionApi),
                        "payloadFetchSolutionApi : " + str(payloadFetchSolutionApi),
                        "Status Code : " + str(responseFetchSolutionApiUrl.status_code),
                        "Response : " + str(responseFetchSolutionApiUrl.text)]
            self.createAPILog(solutionName_for_folder_path, messageArr)
            responseFetchSolutionJson = responseFetchSolutionApiUrl.json()

            if responseFetchSolutionApiUrl.status_code == 200:
                solutionName = responseFetchSolutionJson["result"]["name"]
                print(solutionName,"solutionName")
                # xfile = openpyxl.load_workbook(programdetails)
                # sheet_name = 'Resource Details'.strip()
                # resourceDetailsSheet = xfile[sheet_name]
                resourceDetailsSheet = ProgramGlobalDict.get('Program Resources')
                for solutions in resourceDetailsSheet:
                    if solutionName == solutions.get('Nameofresourcesinprogram'):
                        solutionMainRole = solutions.get('Targetroleattheresourcelevel')
                        solutionRolesArray = solutions.get('Targetedsubroleatresourcelevel')
                        solutionStartDate = solutions.get('Startdateofresource')
                        solutionEndDate = solutions.get('Enddateofresource')
                        return [solutionMainRole,solutionRolesArray, solutionStartDate, solutionEndDate]
            else:
                if responseFetchSolutionApiUrl.status_code in [400, 401, 403, 404, 422]:
                    self.errorVar.append(f"FetchSolutionApiUrl-Client Error {responseFetchSolutionApiUrl.status_code}: {responseFetchSolutionApiUrl.text}")
                elif responseFetchSolutionApiUrl.status_code in [500, 502, 503, 504]:
                    self.errorVar.append(f"FetchSolutionApiUrl-Server Error {responseFetchSolutionApiUrl.status_code}: {responseFetchSolutionApiUrl.text}")
                else:
                    self.errorVar.append(f"FetchSolutionApiUrl-Unexpected Error {responseFetchSolutionApiUrl.status_code}: {responseFetchSolutionApiUrl.text}")
                return False
        except Exception as e:
            messageArr = []
            messageArr.append("Exception caught : " + str(e))
            self.createAPILog(solutionName_for_folder_path, messageArr)
            self.errorVar.append(f"Error occurred: {str(e)}")
            return False
    
    def validate_roles_against_api(self, mainRoles, subRoles, programdetails, parentFolder):
        urlFetchRoleList = userLoginHost + fetchprofessionalRole
        headers = {
            'Content-Type': content_type,
            'tenantId': programdetails.get('TenantID'),
            'X-Channel-id': x_channel_id,
        }
        payload = {}

        response = requests.request("GET", urlFetchRoleList, headers=headers, data=payload)

        messageArr = []
        messageArr = ["Solution Fetch Link.",
                    "URL : " + str(urlFetchRoleList),
                    "Headers: "+ str(headers),
                    "payload : " + str(payload),
                    "Status Code : " + str(response.status_code),
                    "Response : " + str(response.text)]
        self.createAPILog(parentFolder, messageArr)
        
        if response.status_code != 200:
            messageArr.append("Error fetching roles.")
            print("Error fetching roles.")
            return [], []

        role_data = response.json()
        validated_main_role_ids = []
        validated_subrole_ids_list = []

        remaining_subroles = [s.strip() for s in subRoles]

        for role in mainRoles:
            matched = next((r for r in role_data['result']
                            if r.get('externalId', '').strip() == role.strip() or
                            r.get('name', '').strip() == role.strip()), None)

            if matched:
                main_role_id = matched['_id']
                validated_main_role_ids.append(main_role_id)
                subrole_url = f"{userLoginHost}entity-management/v1/entities/subEntityList/{main_role_id}?type=professional_subroles"
                subrole_resp = requests.request("GET",subrole_url, headers=headers,data=payload)
                if subrole_resp.status_code == 200:
                    subroles_data = subrole_resp.json()
                    for item in subroles_data['result']['data']:
                        sub_external_id = item.get('externalId', '').strip()
                        sub_name = item.get('name', '').strip()
                        if sub_external_id in remaining_subroles or sub_name in remaining_subroles:
                            validated_subrole_ids_list.append(item['_id'])
                            if sub_external_id in remaining_subroles:
                                remaining_subroles.remove(sub_external_id)
                            elif sub_name in remaining_subroles:
                                remaining_subroles.remove(sub_name)
                            messageArr.append(f"Subrole '{sub_external_id}' validated under mainRole '{role}'")
                            print(f"Subrole '{sub_external_id}' validated under mainRole '{role}'")
                else:
                    messageArr.append(f"Failed to fetch subroles for mainRole '{role}'")
                    self.errorVar.append(f"Failed to fetch subroles for mainRole '{role}'")
            else:
                messageArr.append(f"MainRole '{role}' not found in API")
                self.errorVar.append(f"MainRole '{role}' not found in API")

        for s in remaining_subroles:
            messageArr.append(f"Subrole '{s}' not found in any of the provided mainRoles.")
            self.errorVar.append(f"Subrole '{s}' not found in any of the provided mainRoles.")
        # Programs.createAPILog(parentFolder, messageArr)
        for err in self.errorVar:
            print(err)
        return validated_main_role_ids, validated_subrole_ids_list 
    
    def prepareProgramSuccessSheet(self, MainFilePath, solutionName_for_folder_path, programFile, solutionExternalId, solutionId,accessToken, programdetails):
        urlFetchSolutionApi = internal_kong_ip + dbfindapi_url
        headerFetchSolutionApi = {
            'Authorization': authorization,
            'X-auth-token': accessToken,
            'Content-Type': content_type,
            'X-Channel-id': x_channel_id,
            'internal-access-token': internal_access_token,
            'tenantId': programdetails.get('TenantID'),
            'orgid': programdetails.get('OrgForAPIs')
            # adminTokenHeaderName: projAdminAccessToken
        }
        payloadFetchSolutionApi = json.dumps({
            "query": {
                "_id": solutionId
            },
            "mongoIdKeys": [
                "_id",
                "solutionId",
                "metaInformation.solutionId"
            ],
            "limit": 10000
        })
        responseFetchSolutionApi = requests.request("POST", url=urlFetchSolutionApi, headers=headerFetchSolutionApi,
                                                data=payloadFetchSolutionApi)
        print(responseFetchSolutionApi.text, "responseFetchSolutionApi")
        responseFetchSolutionJson = responseFetchSolutionApi.json()
        messageArr = []
        messageArr.append("Solution Fetch Link.")
        messageArr.append("solution name : " + responseFetchSolutionJson["result"][0]["name"])
        messageArr.append("solution ExternalId : " + responseFetchSolutionJson["result"][0]["externalId"])
        messageArr.append("Upload status code : " + str(responseFetchSolutionApi.status_code))
        self.createAPILog(solutionName_for_folder_path, messageArr)

        if responseFetchSolutionApi.status_code == 200:
            print('Fetch solution Api Success')
            solutionName = responseFetchSolutionJson["result"][0]["name"]
        urlFetchSolutionLinkApi = internal_kong_ip + fetchlink + solutionId
        print(urlFetchSolutionLinkApi,"urlFetchSolutionLinkApi")
        headerFetchSolutionLinkApi = {
            # 'Authorization': authorization,
            'X-auth-token': accessToken,
            # 'X-Channel-id': x_channel_id
            'internal-access-token': internal_access_token,
            'tenantId': programdetails.get('TenantID'),
            'orgid': programdetails.get('OrgForAPIs')
            #adminTokenHeaderName: projAdminAccessToken
        }
        payloadFetchSolutionLinkApi = {}

        responseFetchSolutionLinkApi = requests.get(url=urlFetchSolutionLinkApi, headers=headerFetchSolutionLinkApi,
                                                    data=payloadFetchSolutionLinkApi)
        
        print(responseFetchSolutionLinkApi.text, "responseFetchSolutionLinkApi")
        messageArr = ["Solution Fetch Link.","solution id : " + solutionId,"solution ExternalId : " + solutionExternalId]
        messageArr.append("Upload status code : " + str(responseFetchSolutionLinkApi.status_code))
        self.createAPILog(solutionName_for_folder_path, messageArr)
        if responseFetchSolutionLinkApi.status_code == 200:
            print(responseFetchSolutionLinkApi.text,"responseFetchSolutionLinkApi")
            print('Fetch solution Link Api Success')
            responseProjectUploadJson = responseFetchSolutionLinkApi.json()
            solutionLink = responseProjectUploadJson["result"]
            solutionLink = ','.join(solutionLink)
            print(solutionLink,"solutionLink")
            messageArr = []
            messageArr.append("Response : " + str(responseFetchSolutionLinkApi.text))
            self.createAPILog(solutionName_for_folder_path, messageArr)
            # Ensure programFile is formatted correctly
            programFileBase = str(programFile).replace(".xlsx", "")
            success_file_path = os.path.join(MainFilePath, programFileBase + '-SuccessSheet.xlsx')

            # Ensure directory exists
            os.makedirs(os.path.dirname(success_file_path), exist_ok=True)

            # Load the workbook correctly
            if os.path.exists(success_file_path):
                xfile = openpyxl.load_workbook(success_file_path)
            else:
                xfile = openpyxl.load_workbook(programFile)

            # Use the correct way to get the sheet
            resourceDetailsSheet = xfile["Resource Details"]  # Instead of get_sheet_by_name

            print(resourceDetailsSheet)

            # Define fill color
            greenFill = PatternFill(start_color='0000FF00', end_color='0000FF00', fill_type='solid')

            # Get row and column count
            rowCountRD = resourceDetailsSheet.max_row
            columnCountRD = resourceDetailsSheet.max_column

            for row in range(3, rowCountRD + 1):
                if str(resourceDetailsSheet["B" + str(row)].value).strip().lower() == "course":
                    resourceDetailsSheet["D1"].value = ""
                    resourceDetailsSheet["E1"].value = ""
                    resourceDetailsSheet['I2'].value = "External id of the resource"
                    resourceDetailsSheet['J2'].value = "link to access the resource/Response"

                    resourceDetailsSheet['I2'].fill = greenFill
                    resourceDetailsSheet['J2'].fill = greenFill
                    resourceDetailsSheet['I' + str(row)].value = solutionExternalId
                    resourceDetailsSheet['J' + str(row)].value = "The course has been successfully mapped to the program"

                    resourceDetailsSheet['I' + str(row)].fill = greenFill
                    resourceDetailsSheet['J' + str(row)].fill = greenFill

                elif str(resourceDetailsSheet["A" + str(row)].value).strip() == solutionName:
                    resourceDetailsSheet["D1"].value = ""
                    resourceDetailsSheet["E1"].value = ""
                    resourceDetailsSheet['I2'].value = "External id of the resource"
                    resourceDetailsSheet['J2'].value = "link to access the resource/Response"

                    resourceDetailsSheet['I2'].fill = greenFill
                    resourceDetailsSheet['J2'].fill = greenFill
                    resourceDetailsSheet['I' + str(row)].value = solutionExternalId
                    resourceDetailsSheet['J' + str(row)].value = solutionLink

                    resourceDetailsSheet['I' + str(row)].fill = greenFill
                    resourceDetailsSheet['J' + str(row)].fill = greenFill

            # Save the file
            xfile.save(success_file_path)
            print("Program success sheet is created")
            return solutionLink
        
    def uploadSurveyQuestionsFromDict(self, MainFilePath, parentFolder, surveyDict, accessToken, surTempExtID, surTempSolID, millisecond, programFile, programdetails,ProgramGlobalDict):
        # Ensure question upload folder exists
        questionFilePath = os.path.join(parentFolder, 'questionUpload')
        os.makedirs(questionFilePath, exist_ok=True)
        csvPath = os.path.join(questionFilePath, 'uploadSheet.csv')

        # Define all CSV fieldnames
        questionUploadFieldnames = [
            'solutionId', 'instanceParentQuestionId', 'hasAParentQuestion',
            'parentQuestionOperator', 'parentQuestionValue', 'parentQuestionId',
            'externalId', 'question0', 'question1', 'tip', 'hint', 'instanceIdentifier',
            'responseType', 'dateFormat', 'autoCapture', 'validation', 'validationIsNumber',
            'validationRegex', 'validationMax', 'validationMin', 'file', 'fileIsRequired',
            'fileUploadType', 'minFileCount', 'maxFileCount', 'caption', 'questionGroup',
            'modeOfCollection', 'accessibility', 'showRemarks', 'rubricLevel', 'isAGeneralQuestion'
        ] + [f'R{i}' for i in range(1, 21)] + [f'R{i}-hint' for i in range(1, 21)] + ['sectionHeader', 'page', 'questionNumber']

        file_exists_ques = os.path.isfile(csvPath)

        # Extract questionsList from surveyDict
        questionsList = surveyDict['questions']
        print(questionsList,"questionsList")

        with open(csvPath, 'a', encoding='utf-8', newline='') as questionUploadFile:
            writer = csv.DictWriter(questionUploadFile, fieldnames=questionUploadFieldnames, lineterminator='\n')
            if not file_exists_ques:
                writer.writeheader()

            quesSeqCnt = 1.0
            questionSeqByEcmArr = []

            for ques in questionsList:
                questionFileObj = {}

                # ---------------- Mandatory Fields ----------------
                questionFileObj['solutionId'] = surTempExtID
                questionFileObj['instanceParentQuestionId'] = (
                    ques.get('instance_parent_question_id', '').strip() + f"_{millisecond}"
                    if ques.get('instance_parent_question_id') else 'NA'
                )

                # Parent question details
                parent_id = ques.get('parent_question_id')
                if parent_id:
                    questionFileObj['hasAParentQuestion'] = 'YES'
                    operator = ques.get('show_when_parent_question_value_is')
                    questionFileObj['parentQuestionOperator'] = '||' if operator == 'or' else operator
                    questionFileObj['parentQuestionValue'] = ques.get('parent_question_value')
                    questionFileObj['parentQuestionId'] = parent_id.strip() + f"_{millisecond}"
                else:
                    questionFileObj['hasAParentQuestion'] = 'NO'
                    questionFileObj['parentQuestionOperator'] = None
                    questionFileObj['parentQuestionValue'] = None
                    questionFileObj['parentQuestionId'] = None

                questionFileObj['externalId'] = ques.get('question_id', '').strip() + f"_{millisecond}"

                if quesSeqCnt == ques.get('question_sequence'):
                    questionSeqByEcmArr.append(questionFileObj['externalId'])
                    quesSeqCnt += 1.0

                # ---------------- Question Texts ----------------
                questionFileObj['question0'] = ques.get('question_language1')
                questionFileObj['question1'] = ques.get('question_language2')
                questionFileObj['tip'] = ques.get('question_tip')
                questionFileObj['hint'] = ques.get('question_hint')
                questionFileObj['instanceIdentifier'] = ques.get('instance_identifier')

                # ---------------- Response Type & Validation ----------------
                qtype = ques.get('question_response_type', '').strip().lower()
                questionFileObj['responseType'] = qtype
                questionFileObj['dateFormat'] = "DD-MM-YYYY" if qtype == 'date' else None
                questionFileObj['autoCapture'] = (
                    'TRUE' if ques.get('date_auto_capture') == 1 else 'false'
                    if qtype == 'date' else None
                )
                questionFileObj['validation'] = 'TRUE' if ques.get('response_required') == 1 else 'FALSE'

                # Number/slider validations
                if qtype in ['number', 'slider']:
                    questionFileObj['validationIsNumber'] = 'TRUE' if qtype == 'number' else None
                    questionFileObj['validationRegex'] = 'isNumber'
                    questionFileObj['validationMax'] = ques.get('max_number_value') or (10000 if qtype == 'number' else 5)
                    questionFileObj['validationMin'] = ques.get('min_number_value') or 0
                else:
                    questionFileObj['validationIsNumber'] = None
                    questionFileObj['validationRegex'] = None
                    questionFileObj['validationMax'] = None
                    questionFileObj['validationMin'] = None

                # ---------------- File Upload ----------------
                if ques.get('file_upload') == 1:
                    questionFileObj['file'] = 'Snapshot'
                    questionFileObj['fileIsRequired'] = 'TRUE'
                    questionFileObj['fileUploadType'] = 'image/jpeg,docx,pdf,ppt'
                    questionFileObj['minFileCount'] = 0
                    questionFileObj['maxFileCount'] = 10
                else:
                    questionFileObj['file'] = 'NA'
                    questionFileObj['fileIsRequired'] = None
                    questionFileObj['fileUploadType'] = None
                    questionFileObj['minFileCount'] = None
                    questionFileObj['maxFileCount'] = None

                # ---------------- Defaults ----------------
                questionFileObj['caption'] = 'FALSE'
                questionFileObj['questionGroup'] = 'A1'
                questionFileObj['modeOfCollection'] = 'onfield'
                questionFileObj['accessibility'] = 'No'
                questionFileObj['showRemarks'] = 'TRUE' if ques.get('show_remarks') == 1 else 'FALSE'
                questionFileObj['rubricLevel'] = None
                questionFileObj['isAGeneralQuestion'] = None
                questionFileObj['sectionHeader'] = ques.get('section_header')
                questionFileObj['page'] = ques.get('page')
                questionFileObj['questionNumber'] = int(ques['question_number']) if ques.get('question_number') else None

                # ---------------- Responses R1-R20 ----------------
                for i in range(1, 21):
                    questionFileObj[f'R{i}'] = ques.get(f'response(R{i})')
                    questionFileObj[f'R{i}-hint'] = ques.get(f'response(R{i})_hint')

                writer.writerow(questionFileObj)

        print(f"Questions CSV created at: {csvPath}")
        try:       
            urlQuestionsUploadApi = internal_kong_ip + questionuploadapiurl
            headerQuestionUploadApi = {
                "internal-access-token": internal_access_token,
                'Authorization': authorization,
                'X-auth-token': accessToken,
                'X-Channel-id': x_channel_id,
                'tenantId': programdetails.get('TenantID'),
                'orgid': programdetails.get('OrgForAPIs')
                # adminTokenHeaderName: adminAccessToken
            }
            filesQuestion = {
                'questions': open(parentFolder + '/questionUpload/uploadSheet.csv', 'rb')
            }
            responseQuestionUploadApi = requests.post(url=urlQuestionsUploadApi,
                                                    headers=headerQuestionUploadApi, files=filesQuestion)
            messageArr = []
            messageArr = ["Question Upload api.",
                        "URL : " + str(urlQuestionsUploadApi),
                        "Headers: "+ str(headerQuestionUploadApi),
                        "file : " + str(filesQuestion),
                        "Status Code : " + str(responseQuestionUploadApi.status_code),
                        "Response : " + str(responseQuestionUploadApi.text)]
            self.createAPILog(parentFolder, messageArr)

            if responseQuestionUploadApi.status_code == 200:
                print('Question upload Success')
                with open(parentFolder + '/questionUpload/uploadInternalIdsSheet.csv', 'w+',encoding='utf-8') as questionRes:
                    questionRes.write(responseQuestionUploadApi.text)
                urlImportSoluTemplate = internal_kong_ip + importsurveysolutiontemplateurl + str(surTempSolID) + "?appName=manage-learn&programId=" + programdetails.get('_id')
                headerImportSoluTemplateApi = {
                    'X-auth-token': accessToken,
                    'X-Channel-id': x_channel_id,
                    'internal-access-token': internal_access_token,
                    'tenantId': programdetails.get('TenantID'),
                    'orgid': programdetails.get('OrgForAPIs')
                    # adminTokenHeaderName: adminAccessToken
                }
                responseImportSoluTemplateApi = requests.post(url=urlImportSoluTemplate,
                                                            headers=headerImportSoluTemplateApi)
                messageArr = []
                messageArr = ["Import Solution Template Api.",
                            "URL : " + str(urlImportSoluTemplate),
                            "Headers: "+ str(headerImportSoluTemplateApi),
                            "Status Code : " + str(responseImportSoluTemplateApi.status_code),
                            "Response : " + str(responseImportSoluTemplateApi.text)]
                self.createAPILog(parentFolder, messageArr)
                if responseImportSoluTemplateApi.status_code == 200:
                    print('Creating Child Success')
                    messageArr = ["********* Creating Child api *********", "URL : " + urlImportSoluTemplate,
                                "Status code : " + str(responseImportSoluTemplateApi.status_code),
                                "Response : " + responseImportSoluTemplateApi.text]
                    # ElevateObservation.createAPILog(parentFolder, messageArr)
                    responseImportSoluTemplateApi = responseImportSoluTemplateApi.json()
                    solutionIdSuc = responseImportSoluTemplateApi["result"]["solutionId"]
                    urlSurveyProgramMapping = internal_kong_ip + importsurveysolutiontoprogramurl + str(solutionIdSuc) + "?programId=" + programdetails.get('_id')
                    headeSurveyProgramMappingApi = {
                        'X-auth-token': accessToken,
                        'X-Channel-id': x_channel_id,
                        'internal-access-token': internal_access_token,
                        'tenantId': programdetails.get('TenantID'),
                        'orgid': programdetails.get('OrgForAPIs')
                        # adminTokenHeaderName: adminAccessToken
                    }
                    responseSurveyProgramMappingApi = requests.post(url=urlSurveyProgramMapping,headers=headeSurveyProgramMappingApi)
                    messageArr = []
                    messageArr = ["Survey Program Mapping Api.",
                                "URL : " + str(urlSurveyProgramMapping),
                                "Headers: "+ str(headeSurveyProgramMappingApi),
                                "Status Code : " + str(responseSurveyProgramMappingApi.status_code),
                                "Response : " + str(responseSurveyProgramMappingApi.text)]
                    self.createAPILog(parentFolder, messageArr)
                    if responseSurveyProgramMappingApi.status_code == 200:
                        print('Program Mapping Success')
                        surveyLink = None
                        solutionIdSuc = None
                        surveyExternalIdSuc = None
                        surveyLink = responseImportSoluTemplateApi["result"]["link"]
                        solutionIdSuc = responseImportSoluTemplateApi["result"]["solutionId"]
                        solutionExtIdSuc = responseImportSoluTemplateApi["result"]["solutionExternalId"]
                        print("Survey Child Id : " + str(solutionExtIdSuc))
                        solutionDetails = self.fetchSolutionDetailsFromProgramSheet(parentFolder, programdetails, solutionIdSuc,
                                                                            accessToken, ProgramGlobalDict)  
                        print(solutionDetails,"solutionDetailssurvey") 
                        scopeRoles = [ solutionDetails[0] ]
                        scopeSubRoles = [ solutionDetails[1] ]
                        verifiedRoles = self.validate_roles_against_api(scopeRoles, scopeSubRoles,programdetails, parentFolder)
                        mainRoleproff = verifiedRoles[0]
                        rolesPGMID = verifiedRoles[1]
                        entities = programdetails.get('entitiesType')
                        entitiesPGMID = entities[1]
                        entitiesType = entities[0]
                        scopeEntities = entitiesPGMID
                        scope = {}
                        entityHierarchy = programdetails.get('entityHierarchy')
                        scope.update(entityHierarchy)
                        scope["professional_subroles"] = rolesPGMID
                        scope["professional_role"] = mainRoleproff
                        bodySolutionUpdate = {
                        "scope": scope
                        }
                        print("scope", bodySolutionUpdate)
                        self.solutionUpdate(parentFolder, accessToken, solutionIdSuc, bodySolutionUpdate, programdetails)
                    
                        if solutionDetails[2]:
                            startDateArr = str(solutionDetails[2]).split("-")
                            bodySolutionUpdate = {
                                "startDate": startDateArr[2] + "-" + startDateArr[1] + "-" + startDateArr[0] + " 00:00:00"}
                            self.solutionUpdate(parentFolder, accessToken, solutionIdSuc, bodySolutionUpdate, programdetails)
                        if solutionDetails[3]:
                            endDateArr = str(solutionDetails[3]).split("-")
                            bodySolutionUpdate = {
                                "endDate": endDateArr[2] + "-" + endDateArr[1] + "-" + endDateArr[0] + " 23:59:59"}
                            self.solutionUpdate(parentFolder, accessToken, solutionIdSuc, bodySolutionUpdate, programdetails)
                        print('Survey Successfully Added')

                        surveySolutionlink = self.prepareProgramSuccessSheet(MainFilePath, parentFolder, programFile, solutionExtIdSuc,
                                            solutionIdSuc, accessToken, programdetails)
                        
                        # surveySolutionlink = "https: Deeplink for survey created successfully."
                        return surveySolutionlink
                    else:
                        print('Program Mapping Failed')
                        if responseSurveyProgramMappingApi.status_code in [400, 401, 403, 404, 422]:
                            self.errorVar.append(f"SurveyProgramMappingApi-Client Error {responseSurveyProgramMappingApi.status_code}: {responseSurveyProgramMappingApi.text}")
                        elif responseSurveyProgramMappingApi.status_code in [500, 502, 503, 504]:
                            self.errorVar.append(f"SurveyProgramMappingApi-Server Error {responseSurveyProgramMappingApi.status_code}: {responseSurveyProgramMappingApi.text}")
                        else:
                            self.errorVar.append(f"SurveyProgramMappingApi-Unexpected Error {responseSurveyProgramMappingApi.status_code}: {responseSurveyProgramMappingApi.text}")
                        messageArr = ["********* Program mapping api *********", "URL : " + urlSurveyProgramMapping,
                                    "Status code : " + str(responseSurveyProgramMappingApi.status_code),
                                    "Response : " + responseSurveyProgramMappingApi.text]
                        # ElevateObservation.createAPILog(parentFolder, messageArr)
                        messageArr.append(f"Error Response: {self.errorVar}")
                else:
                    print('Creating Child API Failed')
                    if responseImportSoluTemplateApi.status_code in [400, 401, 403, 404, 422]:
                        self.errorVar.append(f"ImportSoluTemplateApi-Client Error {responseImportSoluTemplateApi.status_code}: {responseImportSoluTemplateApi.text}")
                    elif responseImportSoluTemplateApi.status_code in [500, 502, 503, 504]:
                        self.errorVar.append(f"ImportSoluTemplateApi-Server Error {responseImportSoluTemplateApi.status_code}: {responseImportSoluTemplateApi.text}")
                    else:
                        self.errorVar.append(f"ImportSoluTemplateApi-Unexpected Error {responseImportSoluTemplateApi.status_code}: {responseImportSoluTemplateApi.text}")
                    messageArr = ["********* Program mapping api *********", "URL : " + urlImportSoluTemplate,
                                "Status code : " + str(responseImportSoluTemplateApi.status_code),
                                "Response : " + responseImportSoluTemplateApi.text]
                    # ElevateObservation.createAPILog(parentFolder, messageArr)
                    messageArr.append(f"Error Response: {self.errorVar}")
            else:
                if responseQuestionUploadApi.status_code in [400, 401, 403, 404, 422]:
                    self.errorVar.append(f"QuestionUploadApi-Client Error {responseQuestionUploadApi.status_code}: {responseQuestionUploadApi.text}")
                elif responseQuestionUploadApi.status_code in [500, 502, 503, 504]:
                    self.errorVar.append(f"QuestionUploadApi-Server Error {responseQuestionUploadApi.status_code}: {responseQuestionUploadApi.text}")
                else:
                    self.errorVar.append(f"QuestionUploadApi-Unexpected Error {responseQuestionUploadApi.status_code}: {responseQuestionUploadApi.text}")
                print('QuestionUploadApi Failed')
                
        except Exception as e:
            messageArr = []
            messageArr.append("Exception caught : " + str(e))
            self.createAPILog(parentFolder, messageArr)
            self.errorVar.append(f"Error occurred: {str(e)}")
            return False

    def CreateSurvey(self, resource, PRName, parentFolder, accessToken, ProgramGlobalDict, programdetails, MainFilePath, programFile):
        if resource.get('typeofSolution') == 3:
            surveySolutionlink = ""
            wbObservation = resource.get("ResourceCre")
            millisecond = int(time.time() * 1000)
            print(wbObservation)
            surveyResp = self.createSurveySolution(parentFolder, wbObservation, accessToken, programdetails)
            if not surveyResp:
                self.errorVar.append('unable to create survey solution.')
                return surveySolutionlink, self.errorVar
            print("survey solution created...")
            surTempExtID = surveyResp[1]
            surTempSolID = surveyResp[0]
            bodySolutionUpdate = {"status": "active", "isDeleted": False}
            if not self.solutionUpdate(parentFolder, accessToken,surTempSolID, bodySolutionUpdate, programdetails):
                self.errorVar.append("Solution Update Failed.")
                return surveySolutionlink, self.errorVar
            surveySolutionlink = self.uploadSurveyQuestionsFromDict(MainFilePath, parentFolder, wbObservation, accessToken, surTempExtID, surTempSolID, millisecond, programFile, programdetails, ProgramGlobalDict)
            if not surveyResp:
                self.errorVar.append('unable to upload survey questions...')
                return surveySolutionlink, self.errorVar
            print(surveySolutionlink,"surveySolutionlink")
            return surveySolutionlink, self.errorVar
